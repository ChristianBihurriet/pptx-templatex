"""Discover themes; presenter notes in the PPTX are the only type registry."""

import hashlib
import io
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation

import plantillas

from .errors import ContentError
from .json_io import read_json
from .metadata import GLOBAL_MARKERS as GLOBAL_MARKERS
from .metadata import MARKER as MARKER
from .metadata import extract_contracts, identifier, keys
from .metadata import text_frames as text_frames


@dataclass(frozen=True)
class Theme:
    name: str
    template: Path
    types: dict
    contracts: dict = field(default_factory=dict)
    source_sha256: str = ""


def theme_roots():
    return list(dict.fromkeys([Path.cwd() / "plantillas", Path(plantillas.__file__).parent]))


def available_themes():
    return sorted({p.parent.name for root in theme_roots()
                   for pattern in ("*/plantilla.pptx", "*/tema.json") for p in root.glob(pattern)})


def load_template(template, name="personal"):
    from .adapter import adapt_content

    template = Path(template)
    try:
        data = template.read_bytes()
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ContentError(f"Tema '{name}': no se puede abrir '{template}': {exc}") from exc
    contracts = extract_contracts(presentation, name)
    types = {
        kind: {"pagina": spec["pagina"], "campos": {
            field: {key: value for key, value in rule.items()
                    if key in {"marcador", "obligatorio", "max_caracteres", "max_lineas"}}
            for field, rule in spec["campos"].items()
        }} for kind, spec in contracts.items()
    }
    theme = Theme(name, template, types, contracts, hashlib.sha256(data).hexdigest())
    for kind, spec in contracts.items():
        try:
            adapt_content({"titulo": "Ejemplo del catálogo", "diapositivas": [spec["ejemplo"]]}, theme)
        except ContentError as exc:
            raise ContentError(f"Tema '{name}', tipo '{kind}', ejemplo inválido: {exc}") from exc
    return theme


def load_theme(name="dam_oscuro"):
    identifier(name, "Nombre de tema")
    directory = next((r / name for r in theme_roots()
                      if (r / name / "plantilla.pptx").is_file() or (r / name / "tema.json").is_file()), None)
    if directory is None:
        raise ContentError(f"Tema '{name}' inexistente. Disponibles: {', '.join(available_themes()) or '(ninguno)'}.")
    filename = "plantilla.pptx"
    config_path = directory / "tema.json"
    if config_path.exists():
        config = read_json(config_path)
        if isinstance(config, dict) and "tipos" in config:
            raise ContentError(
                f"Tema '{name}': migra los contratos de tema.json a notas del PPTX; no se admiten contratos duplicados."
            )
        keys(config, {"version", "plantilla"}, {"version", "plantilla"}, f"Tema '{name}'")
        if type(config["version"]) is not int or config["version"] != 1:
            raise ContentError(f"Tema '{name}': version de configuración debe ser 1.")
        filename = config["plantilla"]
    if not isinstance(filename, str) or not filename.endswith(".pptx"):
        raise ContentError(f"Tema '{name}': falta plantilla .pptx.")
    template = (directory / filename).resolve()
    if template.parent != directory.resolve():
        raise ContentError(f"Tema '{name}': la plantilla debe estar dentro de su directorio.")
    return load_template(template, name)
