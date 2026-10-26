"""Isolate the slide copier from template notes; verify the published package."""

import zipfile
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .errors import ContentError
from .metadata import RESERVED, split_notes


def sanitize_template(source, destination):
    presentation = Presentation(str(source))
    human = {}
    for page, slide in enumerate(presentation.slides, start=1):
        if not slide.has_notes_slide:
            continue
        notes = slide.notes_slide
        frame = notes.notes_text_frame
        if frame is None:
            raise ContentError(f"Página {page}: notas sin marcador de cuerpo.")
        # Only normal notes placeholders are supported. Never discard custom notes silently.
        if any(not s.is_placeholder for s in notes.shapes):
            raise ContentError(f"Página {page}: notas con formas o imágenes personalizadas no admitidas.")
        _, human[page] = split_notes(frame.text, f"Página {page}, notas")
        for rel in list(slide.part.rels.values()):
            if rel.reltype == RT.NOTES_SLIDE:
                slide.part.drop_rel(rel.rId)
    presentation.save(str(destination))
    return human


def restore_notes(path, config, human):
    presentation = Presentation(str(path))
    if len(presentation.slides) != len(config["slides"]):
        raise ContentError("El motor no generó el número de diapositivas esperado.")
    for slide, spec in zip(presentation.slides, config["slides"]):
        text = human.get(spec["src_page"], "")
        if text:
            frame = slide.notes_slide.notes_text_frame
            if frame is None:
                raise ContentError("No se pudieron reconstruir las notas docentes.")
            frame.text = text
    presentation.save(str(path))


def verify_clean_output(path):
    """Fail closed, including metadata hidden in XML runs and notes backlinks."""
    try:
        with zipfile.ZipFile(path) as package:
            for name in package.namelist():
                if name.endswith((".xml", ".rels")):
                    root = ET.fromstring(package.read(name))
                    if RESERVED in "".join(root.itertext()) or RESERVED in package.read(name).decode("utf-8"):
                        raise ContentError(f"Salida bloqueada: metadatos internos en '{name}'.")
        presentation = Presentation(str(path))
        for page, slide in enumerate(presentation.slides, start=1):
            if not slide.has_notes_slide:
                continue
            notes = slide.notes_slide
            if notes.notes_text_frame is None:
                raise ContentError(f"Salida bloqueada: página {page} sin cuerpo de notas.")
            rels = list(notes.part.rels.values())
            back = [r for r in rels if r.reltype == RT.SLIDE and not r.is_external]
            masters = [r for r in rels if r.reltype == RT.NOTES_MASTER and not r.is_external]
            if len(back) != 1 or back[0].target_part is not slide.part or len(masters) != 1:
                raise ContentError(f"Salida bloqueada: relaciones de notas inválidas en página {page}.")
    except (zipfile.BadZipFile, ET.ParseError, KeyError, AttributeError) as exc:
        raise ContentError(f"Salida bloqueada: paquete PowerPoint inválido: {exc}") from exc
