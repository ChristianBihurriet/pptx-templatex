"""Production regression coverage for authoring, extraction and publishing."""

import hashlib
import json
import re
import shutil
import zipfile

import pytest
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from pptgen import adapt_content, generate, load_theme
from pptgen.catalog import catalog_data, catalog_markdown
from pptgen.cli import main
from pptgen.errors import ContentError
from pptgen.metadata import BEGIN, END, RESERVED, format_notes, split_notes
from pptgen.notes import verify_clean_output
from pptgen.themes import load_template


@pytest.fixture
def model(tmp_path):
    source = load_theme().template
    destination = tmp_path / "plantilla.pptx"
    shutil.copyfile(source, destination)
    return destination


def edit_contract(model, mutate, index=0, human=""):
    prs = Presentation(model)
    frame = prs.slides[index].notes_slide.notes_text_frame
    contract, _ = split_notes(frame.text)
    mutate(contract)
    frame.text = format_notes(contract, human)
    prs.save(model)


def test_all_four_types_migrated_with_exact_old_limits(model):
    theme = load_template(model)
    assert list(theme.types)[:4] == ["portada", "pregunta", "definicion", "codigo"]
    assert theme.types["codigo"]["campos"]["codigo"] == {
        "marcador": "codigo", "obligatorio": True, "max_caracteres": 1400, "max_lineas": 18,
    }
    assert theme.types["portada"]["campos"]["subtitulo"]["max_caracteres"] == 180
    assert theme.types["definicion"]["campos"]["contenido"]["max_lineas"] == 8


def test_sole_pptx_theme_discovered_and_documents(tmp_path, monkeypatch, model):
    directory = tmp_path / "plantillas" / "solo_pptx"
    directory.mkdir(parents=True)
    shutil.copyfile(model, directory / "plantilla.pptx")
    monkeypatch.chdir(tmp_path)
    assert main(["--tema", "solo_pptx", "--listar-tipos"]) == 0
    assert main(["--tema", "solo_pptx", "--documentar-tipos", "--salida", "catalogos"]) == 0
    catalog = json.loads((tmp_path / "catalogos/tipos_solo_pptx.json").read_text(encoding="utf-8"))
    assert len(catalog["tipos"]) == len(load_template(model).types)
    assert not (directory / "tema.json").exists()
    text = (tmp_path / "catalogos/tipos_solo_pptx.md").read_text(encoding="utf-8")
    theme = load_theme("solo_pptx")
    for block in re.findall(r"```json\n(.*?)\n```", text, re.S):
        assert adapt_content(json.loads(block), theme)


def test_catalog_matches_contracts_and_reflects_notes_changes(model):
    edit_contract(model, lambda c: c["campos"]["titulo"].update(max_caracteres=101))
    theme = load_template(model)
    data = catalog_data(theme)
    assert data["tipos"]["portada"]["campos"]["titulo"]["max_caracteres"] == 101
    assert "| titulo | string | sí | 101 | 2 |" in catalog_markdown(theme)
    assert data["origen_sha256"] == hashlib.sha256(model.read_bytes()).hexdigest()


def test_cli_modes_are_exclusive():
    with pytest.raises(SystemExit) as exc:
        main(["--documentar-tipos", "--validar"])
    assert exc.value.code == 2


def test_reordering_changes_page_not_type_and_generation_uses_new_page(model, tmp_path):
    prs = Presentation(model)
    first = prs.slides._sldIdLst[0]
    prs.slides._sldIdLst.remove(first)
    prs.slides._sldIdLst.append(first)
    prs.save(model)
    theme = load_template(model)
    assert theme.types["portada"]["pagina"] == len(prs.slides)
    content = {"titulo": "Clase", "diapositivas": [
        {"tipo": "portada", "titulo": "Inicio"}, {"tipo": "portada", "titulo": "Final"},
    ]}
    output = generate(content, theme, "Reordenado", tmp_path)
    prs = Presentation(output)
    assert len(prs.slides) == 2
    assert any(s.has_text_frame and s.text == "Inicio" for s in prs.slides[0].shapes)
    assert any(s.has_text_frame and s.text == "Final" for s in prs.slides[1].shapes)


@pytest.mark.parametrize("mutation,match", [
    (lambda c: c.update(metadatos_version=2), "metadatos_version"),
    (lambda c: c.update(metadatos_version=True), "metadatos_version"),
    (lambda c: c.update(contrato_version=0), "contrato_version"),
    (lambda c: c.update(proposito=[]), "proposito"),
    (lambda c: c.update(restricciones="texto"), "restricciones"),
    (lambda c: c.update(extra=True), "extra"),
    (lambda c: c["campos"]["titulo"].update(tipo_dato="object"), "string"),
    (lambda c: c["campos"]["titulo"].update(max_caracteres=True), "entero"),
    (lambda c: c["campos"]["titulo"].update(obligatorio=1), "booleano"),
    (lambda c: c["campos"]["titulo"].update(marcador="modulo"), "reservado"),
    (lambda c: c["campos"]["titulo"].update(marcador="subtitulo"), "duplicado"),
    (lambda c: c["ejemplo"].pop("titulo"), "ejemplo inválido"),
    (lambda c: c["ejemplo"].update(tipo="codigo"), "mismo tipo"),
])
def test_invalid_contract_reports_page_and_reason(model, mutation, match):
    edit_contract(model, mutation)
    with pytest.raises(ContentError, match=match):
        load_template(model)


@pytest.mark.parametrize("notes", [
    BEGIN, END + BEGIN, BEGIN + "{}" + END + END,
    BEGIN + '{"tipo":"a", "tipo":"b"}' + END,
    BEGIN + '{"tipo":NaN}' + END,
    "[[PPTGEN_METADATA_JSON_V2]]{}[[/PPTGEN_METADATA_JSON_V2]]",
])
def test_bad_notes_rejected_with_context(model, notes):
    prs = Presentation(model)
    prs.slides[0].notes_slide.notes_text_frame.text = notes
    prs.save(model)
    with pytest.raises(ContentError, match="página 1"):
        load_template(model)


def test_duplicate_type_rejected(model):
    prs = Presentation(model)
    prs.slides[1].notes_slide.notes_text_frame.text = prs.slides[0].notes_slide.notes_text_frame.text
    prs.save(model)
    with pytest.raises(ContentError, match="duplicado"):
        load_template(model)


def test_missing_metadata_not_filled_from_legacy_config(model, tmp_path, monkeypatch):
    directory = tmp_path / "plantillas" / "antiguo"
    directory.mkdir(parents=True)
    shutil.copyfile(model, directory / "plantilla.pptx")
    (directory / "tema.json").write_text('{"version":1,"plantilla":"plantilla.pptx","tipos":{}}')
    monkeypatch.chdir(tmp_path)
    with pytest.raises(ContentError, match="migra"):
        load_theme("antiguo")


def test_missing_block_is_error_but_static_unannotated_page_is_allowed(model):
    prs = Presentation(model)
    prs.slides[0].notes_slide.notes_text_frame.text = "Sin contrato"
    prs.save(model)
    with pytest.raises(ContentError, match="sin contrato"):
        load_template(model)
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            shape.text = "Texto fijo"
    prs.save(model)
    assert len(load_template(model).types) == len(prs.slides) - 1


def test_no_notes_body_is_clear_error(model):
    prs = Presentation(model)
    placeholder = prs.slides[0].notes_slide.notes_placeholder
    placeholder.element.getparent().remove(placeholder.element)
    prs.save(model)
    with pytest.raises(ContentError, match="cuerpo"):
        load_template(model)


def test_notes_are_clean_valid_and_backlink_to_each_repeated_slide(model, tmp_path):
    edit_contract(model, lambda c: None, human="Consejo docente: usar un ejemplo cotidiano.")
    theme = load_template(model)
    before = model.read_bytes()
    content = {"titulo": "Clase", "diapositivas": [
        {"tipo": "portada", "titulo": "Uno"}, {"tipo": "portada", "titulo": "Dos"},
        {"tipo": "codigo", "titulo": "Java", "codigo": "int n = 2 * 3 * 4;"},
    ]}
    output = generate(content, theme, "Notas", tmp_path)
    assert model.read_bytes() == before
    verify_clean_output(output)
    with zipfile.ZipFile(output) as package:
        for name in package.namelist():
            if name.endswith((".xml", ".rels")):
                data = package.read(name)
                assert RESERVED.encode() not in data
                assert b"contrato_version" not in data
                assert b"proposito" not in data
    prs = Presentation(output)
    for slide in list(prs.slides)[:2]:
        notes = slide.notes_slide
        assert notes.notes_text_frame.text == "Consejo docente: usar un ejemplo cotidiano."
        assert notes.part.part_related_by(RT.SLIDE) is slide.part
        assert notes.part.part_related_by(RT.NOTES_MASTER)
    assert not prs.slides[2].has_notes_slide
    assert not list(tmp_path.glob(".pptgen-*"))


def test_generation_reloads_contract_instead_of_using_stale_cached_theme(model, tmp_path):
    theme = load_template(model)
    edit_contract(model, lambda c: c["campos"]["titulo"].update(max_caracteres=40))
    content = {"titulo": "Clase", "diapositivas": [{"tipo": "portada", "titulo": "x" * 50}]}
    with pytest.raises(ContentError, match="40 caracteres"):
        generate(content, theme, "No generar", tmp_path)
    assert not (tmp_path / "No generar.pptx").exists()


def test_verification_blocks_metadata_outside_notes_before_publication(model, tmp_path):
    prs = Presentation(model)
    prs.core_properties.comments = BEGIN
    prs.save(model)
    with pytest.raises(ContentError, match="Salida bloqueada"):
        generate({"titulo": "Clase", "diapositivas": [{"tipo": "portada", "titulo": "A"}]},
                 load_template(model), "No publicar", tmp_path)
    assert not (tmp_path / "No publicar.pptx").exists()
    assert not list(tmp_path.glob(".pptgen-*"))


def test_invalid_notes_relationship_is_detected(model):
    prs = Presentation(model)
    notes = prs.slides[0].notes_slide
    rel = next(r for r in notes.part.rels.values() if r.reltype == RT.SLIDE)
    notes.part.drop_rel(rel.rId)
    # Remove metadata too: this test targets the graph, not the marker scanner.
    for slide in prs.slides:
        slide.notes_slide.notes_text_frame.text = "Notas"
    prs.save(model)
    with pytest.raises(ContentError, match="relaciones de notas"):
        verify_clean_output(model)


def test_optional_filename_pointer_prevents_path_escape(tmp_path, monkeypatch):
    directory = tmp_path / "plantillas" / "escape"
    directory.mkdir(parents=True)
    (directory / "tema.json").write_text('{"version":1,"plantilla":"../fuera.pptx"}')
    monkeypatch.chdir(tmp_path)
    with pytest.raises(ContentError, match="directorio"):
        load_theme("escape")
