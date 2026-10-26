"""Contract, theme compatibility, generation and CLI regression tests."""

import copy
import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

import plantillas
from pptgen import adapt_content, generate, load_theme
from pptgen.cli import main
from pptgen.errors import ContentError
from pptgen.json_io import read_json
from pptgen.metadata import format_notes, split_notes


@pytest.fixture
def theme():
    return load_theme("dam_oscuro")


@pytest.fixture
def content():
    return {
        "titulo": "Herencia", "modulo": "Programación", "curso": "1DAM",
        "diapositivas": [
            {"tipo": "portada", "titulo": "Inicio"},
            {"tipo": "pregunta", "titulo": "¿Por qué?", "contenido": "Contexto real"},
            {"tipo": "definicion", "titulo": "Concepto", "contenido": "Teoría"},
            {"tipo": "codigo", "titulo": "Java", "codigo": "class Perro extends Animal {}"},
            {"tipo": "pregunta", "titulo": "¿Cuándo?", "contenido": "Otro contexto"},
        ],
    }


@pytest.fixture
def local_theme(tmp_path, monkeypatch):
    root = tmp_path / "plantillas" / "personal"
    shutil.copytree(Path(plantillas.__file__).parent / "dam_oscuro", root)
    monkeypatch.chdir(tmp_path)
    return root


def local_template(root):
    return root / read_json(root / "tema.json")["plantilla"]


def update_contract(root, kind, mutate):
    path = local_template(root)
    presentation = Presentation(path)
    for slide in presentation.slides:
        frame = slide.notes_slide.notes_text_frame
        contract, human = split_notes(frame.text)
        if contract and contract["tipo"] == kind:
            mutate(contract)
            frame.text = format_notes(contract, human)
            break
    presentation.save(path)


def test_adapter_maps_types_repeated_pages_and_optional_fields(content, theme):
    original = copy.deepcopy(content)
    config = adapt_content(content, theme)
    assert [s["src_page"] for s in config["slides"]] == [1, 2, 3, 4, 2]
    assert config["slides"][0]["replace_texts"] == {
        "titulo_clase": "Herencia", "modulo": "Programación", "curso": "1DAM",
        "titulo": "Inicio", "subtitulo": "",
    }
    assert config["slides"][3]["replace_texts"]["lenguaje"] == ""
    assert content == original


@pytest.mark.parametrize("bad,match", [
    ([], "objeto"),
    ({}, "titulo"),
    ({"titulo": " "}, "titulo"),
    ({"titulo": "A", "diapositivas": []}, "diapositivas"),
    ({"titulo": "A", "diapositivas": {}}, "diapositivas"),
    ({"titulo": "A", "diapositivas": [None]}, "Diapositiva 1"),
    ({"titulo": "A", "diapositivas": [{}]}, "tipo"),
    ({"titulo": "A", "diapositivas": [{"tipo": "tipo_inexistente"}]}, "desconocido"),
])
def test_invalid_document(theme, bad, match):
    with pytest.raises(ContentError, match=match):
        adapt_content(bad, theme)


@pytest.mark.parametrize("key,value,match", [
    ("titulo", None, "titulo.*texto"),
    ("titulo", " ", "titulo.*obligatorio"),
    ("contenido", 123, "contenido.*texto"),
    ("contenido", ["texto"], "contenido.*texto"),
    ("contenido", "x" * 501, "contenido.*500 caracteres"),
    ("contenido", "\n".join(["x"] * 9), "contenido.*8 líneas"),
    ("contenido", "a\x0bb", "contenido.*control"),
    ("contenido", "Ejemplo {{ img:foto }}", "contenido.*reservada"),
    ("color", "rojo", "desconocidos.*color"),
])
def test_field_errors_identify_slide(content, theme, key, value, match):
    content["diapositivas"][1][key] = value
    with pytest.raises(ContentError, match=f"Diapositiva 2.*{match}"):
        adapt_content(content, theme)


def test_missing_required_field(content, theme):
    del content["diapositivas"][1]["contenido"]
    with pytest.raises(ContentError, match="Diapositiva 2.*contenido.*obligatorio"):
        adapt_content(content, theme)


def test_unknown_root_field(content, theme):
    content["coordenadas"] = {}
    with pytest.raises(ContentError, match="coordenadas"):
        adapt_content(content, theme)


def test_unknown_theme():
    with pytest.raises(ContentError, match="inexistente"):
        load_theme("no_existe")
    with pytest.raises(ContentError, match="no válido"):
        load_theme("../dam_oscuro")


@pytest.mark.parametrize("change,match", [
    (lambda c: c.update(metadatos_version=99), "metadatos_version"),
    (lambda c: c.update(contrato_version=True), "contrato_version"),
    (lambda c: c.update(pagina=99), "desconocidos.*pagina"),
    (lambda c: c["campos"]["titulo"].update(marcador="ausente"), "marcadores"),
    (lambda c: c["campos"]["titulo"].update(max_lineas=0), "max_lineas"),
    (lambda c: c["campos"]["titulo"].update(obligatorio="sí"), "booleano"),
])
def test_incompatible_theme(local_theme, change, match):
    update_contract(local_theme, "portada", change)
    with pytest.raises(ContentError, match=match):
        load_theme("personal")


def test_custom_type_without_python_changes(local_theme):
    def change(contract):
        contract["tipo"] = "resumen_personal"
        contract["ejemplo"]["tipo"] = "resumen_personal"
    update_contract(local_theme, "definicion", change)
    config = adapt_content({"titulo": "Clase", "diapositivas": [
        {"tipo": "resumen_personal", "titulo": "Recuerda", "contenido": "Idea clave"},
    ]}, load_theme("personal"))
    assert config["slides"][0]["src_page"] == 3


def test_custom_field_mapping(local_theme):
    def change(config):
        fields = config["campos"]
        fields["enunciado"] = fields.pop("contenido")
        config["ejemplo"]["enunciado"] = config["ejemplo"].pop("contenido")
    update_contract(local_theme, "pregunta", change)
    config = adapt_content({"titulo": "Clase", "diapositivas": [
        {"tipo": "pregunta", "titulo": "¿Por qué?", "enunciado": "Problema"},
    ]}, load_theme("personal"))
    assert config["slides"][0]["replace_texts"]["contenido"] == "Problema"


def test_template_markers_validated(local_theme):
    path = local_template(local_theme)
    presentation = Presentation(path)
    presentation.slides[0].shapes[1].text = "{{ marcador_desconocido }}"
    presentation.save(path)
    with pytest.raises(ContentError, match="marcadores.*desconocido"):
        load_theme("personal")


def test_generate_editable_slides_and_preserve_literal_code(tmp_path, content, theme):
    code = 'class Main {\n    int x = 2 * 3 * 4; // **literal**\n    String s = "{{ hola }}";\n}'
    content["diapositivas"][3]["codigo"] = code
    output = generate(content, theme, "Herencia en Java", tmp_path / "salida")
    presentation = Presentation(output)
    assert len(presentation.slides) == 5
    texts = [[s.text for s in slide.shapes if s.has_text_frame] for slide in presentation.slides]
    assert "Contexto real" in texts[1]
    assert "Otro contexto" in texts[4]
    assert code in texts[3]
    assert "{{ subtitulo }}" not in " ".join(texts[0])
    code_shape = next(s for s in presentation.slides[3].shapes if s.has_text_frame and s.text == code)
    assert code_shape.text_frame.paragraphs[0].runs[0].font.name == "Consolas"
    assert output.name == "Herencia en Java.pptx"
    with pytest.raises(ContentError, match="ya existe"):
        generate(content, theme, "Herencia en Java", output.parent)


@pytest.mark.parametrize("name", ["../fuera", "a/b", "CON", "lpt1", "a:", "", "nombre."])
def test_invalid_output_names(tmp_path, content, theme, name):
    with pytest.raises(ContentError, match="Nombre"):
        generate(content, theme, name, tmp_path)


def test_cli_validate_lists_and_generate(tmp_path, content, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "clase.json"
    path.write_text(json.dumps(content), encoding="utf-8")
    assert main(["--listar-temas"]) == 0
    assert "dam_oscuro" in capsys.readouterr().out
    assert main(["--listar-tipos"]) == 0
    assert "codigo" in capsys.readouterr().out
    assert main([str(path), "--validar"]) == 0
    assert not (tmp_path / "salida").exists()
    assert main([str(path)]) == 0
    assert (tmp_path / "salida" / "clase.pptx").is_file()
    assert main([str(path), "--tema", "inexistente"]) == 1
    assert "inexistente" in capsys.readouterr().err
    assert main([str(path), "--nombre", "Nombre explícito", "--salida", str(tmp_path / "otra")]) == 0
    assert (tmp_path / "otra" / "Nombre explícito.pptx").is_file()


@pytest.mark.parametrize("data,match", [
    ('{"titulo": "a", "titulo": "b"}', "duplicada"),
    ('{"titulo": NaN}', "Constante"),
    ('{"titulo": ', "JSON"),
])
def test_json_errors(tmp_path, data, match):
    path = tmp_path / "bad.json"
    path.write_text(data, encoding="utf-8")
    with pytest.raises(ContentError, match=match):
        read_json(path)
    assert main([str(path), "--validar"]) == 1


def test_bundled_example_generates(tmp_path, theme):
    content_path = Path(__file__).resolve().parents[1] / "contenidos/programacion/herencia.json"
    output = generate(read_json(content_path), theme, "Ejemplo", tmp_path)
    assert len(Presentation(output).slides) == 8
