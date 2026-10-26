"""Acceptance tests for the editable DAM layout library, not a type registry."""

import copy
import json
import posixpath
import re
import zipfile
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE

from pptgen import adapt_content, generate, load_theme
from pptgen.catalog import catalog_data, catalog_markdown
from pptgen.errors import ContentError
from pptgen.json_io import read_json
from pptgen.metadata import RESERVED, slide_markers, text_frames
from pptgen.notes import verify_clean_output
from pptgen.themes import load_template

ROOT = Path(__file__).resolve().parents[1]
KINDS = (
    "portada", "pregunta", "definicion", "codigo", "explicacion", "problema",
    "codigo_explicado", "codigo_resultado", "comparacion", "antes_despues", "error_correccion",
    "pasos", "tabla", "diagrama_flujo", "relaciones_clases", "caso_practico", "advertencia",
    "resumen", "terminal", "concepto_visual",
)


@pytest.fixture(scope="module")
def library():
    return load_theme()


@pytest.fixture(scope="module")
def demonstration(library, tmp_path_factory):
    content = read_json(ROOT / "contenidos/demostracion/biblioteca_dam.json")
    output = generate(content, library, "Biblioteca", tmp_path_factory.mktemp("biblioteca"))
    return content, output


def test_library_contains_exactly_the_requested_types(library):
    assert tuple(library.types) == KINDS
    assert len(library.types) == len(set(library.types)) == 20
    prs = Presentation(library.template)
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < .001
    for kind, contract in library.contracts.items():
        slide = prs.slides[contract["pagina"] - 1]
        assert {r["marcador"] for r in contract["campos"].values()} <= slide_markers(slide, kind)
        if kind not in KINDS[:4]:
            assert all(r.get("descripcion", "").strip() for r in contract["campos"].values())
            assert contract["restricciones"]
            assert all("max_caracteres" in r and "max_lineas" in r for r in contract["campos"].values())


def test_original_four_slides_and_contracts_are_unchanged(library):
    baseline = load_template(library.template.with_name("docente.pptx"))
    old = Presentation(baseline.template)
    current = Presentation(library.template)
    assert {kind: library.contracts[kind] for kind in baseline.contracts} == baseline.contracts
    for before, after in zip(old.slides, list(current.slides)[:4]):
        assert before.element.xml == after.element.xml
        assert before.notes_slide.element.xml == after.notes_slide.element.xml
    adapt_content(read_json(ROOT / "contenidos/programacion/herencia.json"), library)


@pytest.mark.parametrize("kind", KINDS)
def test_each_contract_example_and_its_limits(library, kind):
    contract = library.contracts[kind]
    example = contract["ejemplo"]
    assert adapt_content({"titulo": "Clase", "diapositivas": [example]}, library)
    for name, rule in contract["campos"].items():
        for value in ("x" * (rule["max_caracteres"] + 1), "\n".join(["x"] * (rule["max_lineas"] + 1))):
            invalid = {**example, name: value}
            with pytest.raises(ContentError, match=name):
                adapt_content({"titulo": "Clase", "diapositivas": [invalid]}, library)
        omitted = {k: v for k, v in example.items() if k != name}
        if rule.get("obligatorio", False):
            with pytest.raises(ContentError, match=name):
                adapt_content({"titulo": "Clase", "diapositivas": [omitted]}, library)
        else:
            config = adapt_content({"titulo": "Clase", "diapositivas": [omitted]}, library)
            assert config["slides"][0]["replace_texts"][rule["marcador"]] == ""


def test_demonstration_covers_every_type_and_replaces_every_field(library, demonstration):
    content, output = demonstration
    assert [s["tipo"] for s in content["diapositivas"]] == list(KINDS)
    prs = Presentation(output)
    assert len(prs.slides) == 20
    for spec, slide in zip(content["diapositivas"], prs.slides):
        actual = [frame.text for frame in text_frames(slide.shapes)]
        for name in library.contracts[spec["tipo"]]["campos"]:
            assert spec[name].strip(), (spec["tipo"], name)
            assert spec[name] in actual, (spec["tipo"], name)
        assert not re.search(r"\{\{\s*[a-z_]+\s*\}\}", " ".join(actual))
    verify_clean_output(output)
    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
        for name in package.namelist():
            if name.endswith((".xml", ".rels")):
                xml = package.read(name)
                ET.fromstring(xml)
                assert RESERVED.encode() not in xml
                assert b"contrato_version" not in xml


def test_native_table_and_diagram_objects_survive_generation_and_editing(library, demonstration, tmp_path):
    _, output = demonstration
    prs = Presentation(output)
    pages = {kind: prs.slides[index] for index, kind in enumerate(KINDS)}
    tables = [s.table for s in pages["tabla"].shapes if s.has_table]
    assert len(tables) == 1
    assert (len(tables[0].rows), len(tables[0].columns)) == (5, 3)
    tables[0].cell(1, 0).text = "Celda editada"
    for kind, expected in [("diagrama_flujo", 3), ("relaciones_clases", 2)]:
        slide = pages[kind]
        shapes = {s.shape_id: s for s in slide.shapes}
        connectors = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
        assert len(connectors) == expected
        for connector in connectors:
            assert connector.element.xpath(".//a:tailEnd[@type='triangle']")
            start = int(connector.element.xpath(".//a:stCxn")[0].get("id"))
            end = int(connector.element.xpath(".//a:endCxn")[0].get("id"))
            assert start in shapes and end in shapes and start != end
            if kind == "diagrama_flujo":
                assert shapes[start].left < shapes[end].left
            else:
                assert shapes[start].top > shapes[end].top
            shapes[start].fill.fore_color.rgb = shapes[end].fill.fore_color.rgb
    title = next(s for s in pages["comparacion"].shapes if s.has_text_frame and s.name == "titulo")
    title.text = "Título editado"
    edited = tmp_path / "editable.pptx"
    prs.save(edited)
    reloaded = Presentation(edited)
    assert any(s.has_table and s.table.cell(1, 0).text == "Celda editada"
               for s in reloaded.slides[KINDS.index("tabla")].shapes)
    assert any(s.has_text_frame and s.text == "Título editado"
               for s in reloaded.slides[KINDS.index("comparacion")].shapes)
    verify_clean_output(edited)


def test_package_relationship_targets_are_resolvable(library, demonstration):
    for path in (library.template, demonstration[1]):
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            for name in names:
                if not name.endswith(".rels"):
                    continue
                # foo/_rels/bar.xml.rels resolves targets relative to foo/bar.xml.
                directory = posixpath.dirname(posixpath.dirname(name))
                for rel in ET.fromstring(package.read(name)):
                    if rel.get("TargetMode") == "External":
                        continue
                    target = unquote(rel.attrib["Target"])
                    resolved = posixpath.normpath(posixpath.join(directory, target)).lstrip("/")
                    assert resolved in names, (path.name, name, target)
            content_types = ET.fromstring(package.read("[Content_Types].xml"))
            for item in content_types:
                if "PartName" in item.attrib:
                    assert unquote(item.attrib["PartName"]).lstrip("/") in names


def test_new_layouts_are_ungrouped_with_fixed_fonts_and_in_bounds(library):
    prs = Presentation(library.template)
    for slide in list(prs.slides)[4:]:
        for shape in slide.shapes:
            assert shape.shape_type != MSO_SHAPE_TYPE.GROUP
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
            for frame in text_frames([shape]):
                if frame.text.strip():
                    assert frame.auto_size == MSO_AUTO_SIZE.NONE
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if "{{" in run.text and run.font.name == "Consolas":
                            assert run.font.size.pt >= 18


def test_reversing_all_pages_preserves_semantic_generation(library, tmp_path):
    prs = Presentation(library.template)
    original_order = list(prs.slides._sldIdLst)
    for element in original_order:
        prs.slides._sldIdLst.remove(element)
    for element in reversed(original_order):
        prs.slides._sldIdLst.append(element)
    path = tmp_path / "invertido.pptx"
    prs.save(path)
    theme = load_template(path)
    content = read_json(ROOT / "contenidos/demostracion/biblioteca_dam.json")
    for kind in KINDS:
        assert theme.types[kind]["pagina"] == 21 - library.types[kind]["pagina"]
    output = generate(content, theme, "Resultado invertido", tmp_path)
    for spec, slide in zip(content["diapositivas"], Presentation(output).slides):
        assert spec["titulo"] in [f.text for f in text_frames(slide.shapes)]
    verify_clean_output(output)


def test_stress_example_reaches_both_limits_and_generates(library, tmp_path):
    content = read_json(ROOT / "contenidos/demostracion/biblioteca_limites.json")
    for spec in content["diapositivas"]:
        for name, rule in library.contracts[spec["tipo"]]["campos"].items():
            assert len(spec[name]) == rule["max_caracteres"]
            assert len(spec[name].split("\n")) == rule["max_lineas"]
    output = generate(content, library, "Límites", tmp_path)
    assert len(Presentation(output).slides) == 20
    verify_clean_output(output)


def test_catalogs_are_fresh_and_every_markdown_example_is_valid(library):
    assert (ROOT / "prompts/tipos_dam_oscuro.md").read_text(encoding="utf-8") == catalog_markdown(library)
    assert read_json(ROOT / "prompts/tipos_dam_oscuro.json") == catalog_data(library)
    text = catalog_markdown(library)
    samples = re.findall(r"```json\n(.*?)\n```", text, re.S)
    assert len(samples) == 20
    for sample in samples:
        assert adapt_content(json.loads(sample), library)


def test_fixed_steps_accept_three_but_no_dynamic_objects(library):
    spec = copy.deepcopy(library.contracts["pasos"]["ejemplo"])
    del spec["paso_4"], spec["paso_5"]
    config = adapt_content({"titulo": "Clase", "diapositivas": [spec]}, library)
    assert config["slides"][0]["replace_texts"]["paso_4"] == ""
    assert config["slides"][0]["replace_texts"]["paso_5"] == ""
    for kind, field, value in [("tabla", "fila_1_col_1", ["una", "lista"]),
                               ("diagrama_flujo", "etapa_1", {"texto": "caja"})]:
        spec = {**library.contracts[kind]["ejemplo"], field: value}
        with pytest.raises(ContentError, match=field):
            adapt_content({"titulo": "Clase", "diapositivas": [spec]}, library)
