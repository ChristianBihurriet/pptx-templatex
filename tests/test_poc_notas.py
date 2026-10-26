"""Experimental notes POC only; no production theme-loader migration."""

import json
from pathlib import Path

import pytest
from pptx import Presentation

from tools.poc_notas_plantilla import BEGIN, END, extract_catalog, run, split_notes


@pytest.fixture(scope="module")
def pilot(tmp_path_factory):
    output = tmp_path_factory.mktemp("notes") / "pilot"
    report = run(output)
    return output, report


def test_pilot_roundtrip_copy_sanitation_and_original_unchanged(pilot):
    output, report = pilot
    assert report["notas_escritas_y_releidas"]
    assert report["salida_sin_sanear_contiene_metadatos"]
    assert all(note["relaciones"] == [] for note in report["notas_salida_sin_sanear"])
    assert not report["salida_limpia_contiene_metadatos"]
    assert not report["salida_limpia_contiene_contrato"]
    assert report["subtitulo_opcional_vacio"] and report["plantilla_original_intacta"]
    assert (output / "catalogo_ia.md").is_file()
    catalog = json.loads((output / "catalogo.json").read_text(encoding="utf-8"))
    assert list(catalog["tipos"]) == ["portada"]
    final = Presentation(str(output / "presentacion_limpia.pptx"))
    assert len(final.slides) == 2
    for slide in final.slides:
        assert slide.notes_slide.notes_text_frame.text.startswith("Recordatorio docente:")
        assert BEGIN not in slide.notes_slide.notes_text_frame.text
    assert any(s.has_text_frame and s.text == "Herencia en Java" for s in final.slides[0].shapes)


def test_type_identity_survives_reordering(pilot, tmp_path):
    output, _ = pilot
    prs = Presentation(str(output / "plantilla_con_metadatos.pptx"))
    first_id = prs.slides._sldIdLst[0]
    prs.slides._sldIdLst.remove(first_id)
    prs.slides._sldIdLst.append(first_id)
    path = tmp_path / "reordered.pptx"
    prs.save(str(path))
    contract = extract_catalog(path)["tipos"]["portada"]
    assert contract["pagina"] == len(prs.slides)
    assert contract["contrato_version"] == 1


def test_duplicate_type_rejected(pilot, tmp_path):
    output, _ = pilot
    prs = Presentation(str(output / "plantilla_con_metadatos.pptx"))
    prs.slides[1].notes_slide.notes_text_frame.text = prs.slides[0].notes_slide.notes_text_frame.text
    path = tmp_path / "duplicate.pptx"
    prs.save(str(path))
    with pytest.raises(ValueError, match="duplicado"):
        extract_catalog(path)


def test_marker_mismatch_rejected(pilot, tmp_path):
    output, _ = pilot
    prs = Presentation(str(output / "plantilla_con_metadatos.pptx"))
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and shape.text == "{{ titulo }}":
            shape.text = "{{ otro }}"
    path = tmp_path / "mismatch.pptx"
    prs.save(str(path))
    with pytest.raises(ValueError, match="no coinciden"):
        extract_catalog(path)


@pytest.mark.parametrize("notes", [BEGIN, END, BEGIN + "{}" + END + END, BEGIN + "not json" + END])
def test_invalid_metadata_is_not_silently_ignored(notes):
    with pytest.raises(ValueError):
        split_notes(notes)


def test_unannotated_notes_are_preserved():
    assert split_notes("Notas del profesor") == (None, "Notas del profesor")


def test_extraction_does_not_create_notes_on_unannotated_pages(pilot):
    output, _ = pilot
    path = Path(output) / "plantilla_con_metadatos.pptx"
    before = path.read_bytes()
    extract_catalog(path)
    assert path.read_bytes() == before
