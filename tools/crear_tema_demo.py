"""Recreate the local demonstration theme: python tools/crear_tema_demo.py.

All visual decisions belong here or in PowerPoint, never in lesson JSON.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from pptgen.metadata import format_notes
from pptgen.themes import load_theme


def create_demo(destination):
    # Contracts are copied from the authoritative PPTX, never duplicated here.
    theme = load_theme("dam_oscuro")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    def textbox(slide, text, x, y, width, height, size, color, font="Aptos", bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if font == "Consolas":
            # 18 literal lines fit the fixed code frame without shrinking.
            paragraph.line_spacing = Pt(15)
            run.font.size = Pt(14)

    for kind in ("portada", "pregunta", "definicion", "codigo"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("101827")
        textbox(slide, "DAM / " + kind.upper(), .7, .35, 11.9, .4, 13, "50D5CB", bold=True)
        textbox(slide, "{{ titulo }}", .7, 1.1, 11.9, 1.35, 34, "F3F6FB", bold=True)
        if kind == "portada":
            textbox(slide, "{{ subtitulo }}", .7, 3, 11.9, 2, 26, "CED7E4")
        elif kind == "codigo":
            textbox(slide, "{{ lenguaje }}", .7, 2.35, 11.9, .4, 14, "50D5CB")
            textbox(slide, "{{ codigo }}", .7, 2.85, 11.9, 3.95, 16, "E7EDF6", font="Consolas")
        else:
            textbox(slide, "{{ contenido }}", .7, 2.9, 11.9, 3.65, 26, "CED7E4")
        textbox(slide, "{{ modulo }}  |  {{ curso }}", .7, 7, 11.9, .3, 11, "9AA9BE")
        contract = {key: value for key, value in theme.contracts[kind].items() if key != "pagina"}
        slide.notes_slide.notes_text_frame.text = format_notes(contract)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(destination))


if __name__ == "__main__":
    target = load_theme("dam_oscuro").template
    create_demo(target)
    print(f"Plantilla creada: {target}")
