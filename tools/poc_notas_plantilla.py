"""Isolated experiment, not the production theme loader or CLI.

Run from the repository: python -m tools.poc_notas_plantilla
Artifacts: salida/poc_notas/ (refuses to overwrite an existing directory).
"""

import argparse
import hashlib
import json
import re
import zipfile
from importlib.metadata import version
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from pptgen.adapter import adapt_content
from pptgen.themes import GLOBAL_MARKERS, MARKER, Theme, load_theme, text_frames
from pptx_templatex import TemplateEngine

BEGIN = "[[PPTGEN_METADATA_JSON_V1]]"
END = "[[/PPTGEN_METADATA_JSON_V1]]"


def split_notes(text):
    """Return contract and human notes; malformed marked blocks are errors."""
    if BEGIN not in text and END not in text:
        return None, text
    if text.count(BEGIN) != 1 or text.count(END) != 1:
        raise ValueError("Se requiere exactamente un bloque de metadatos completo por página.")
    before, body = text.split(BEGIN)
    payload, after = body.split(END)
    contract = json.loads(payload)
    if not isinstance(contract, dict):
        raise ValueError("Los metadatos deben ser un objeto JSON.")
    return contract, (before.rstrip() + "\n" + after.lstrip()).strip()


def extract_catalog(template):
    """POC: only the PPTX is read; no tema.json lookup during extraction."""
    presentation = Presentation(str(template))
    catalog = {"catalogo_version": 1, "tipos": {}}
    for page, slide in enumerate(presentation.slides, start=1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            raise ValueError(f"Página {page}: notas sin cuadro de texto.")
        contract, _ = split_notes(frame.text)
        if contract is None:
            continue  # Deliberately partial pilot: unannotated pages are not types.
        if contract.get("metadatos_version") != 1 or contract.get("contrato_version") != 1:
            raise ValueError(f"Página {page}: versión no admitida en esta prueba.")
        kind = contract["tipo"]
        if not re.fullmatch(r"[a-z][a-z0-9_]*", kind) or kind in catalog["tipos"]:
            raise ValueError(f"Página {page}: identificador inválido o tipo duplicado: {kind}.")
        fields = contract["campos"]
        markers = [rule["marcador"] for rule in fields.values()]
        if len(set(markers)) != len(markers) or set(markers) & GLOBAL_MARKERS:
            raise ValueError(f"Página {page}: marcadores duplicados o reservados.")
        found = set()
        for shape in slide.shapes:
            if hasattr(shape, "shapes"):
                raise ValueError(f"Página {page}: formas agrupadas no admitidas.")
        for text_frame in text_frames(slide.shapes):
            for paragraph in text_frame.paragraphs:
                found.update(MARKER.findall(paragraph.text))
                residue = MARKER.sub("", paragraph.text)
                if "{{" in residue or "}}" in residue:
                    raise ValueError(f"Página {page}: marcador mal formado.")
        if set(markers) - found or found - set(markers) - GLOBAL_MARKERS:
            raise ValueError(f"Página {page}: campos y marcadores no coinciden.")
        catalog["tipos"][kind] = {**contract, "pagina": page}
    catalog["origen_sha256"] = hashlib.sha256(Path(template).read_bytes()).hexdigest()
    return catalog


def adapter_theme(catalog, template):
    """Normalize the proposed metadata to the CURRENT adapter's existing model."""
    types = {}
    for kind, contract in catalog["tipos"].items():
        fields = {}
        for field, rule in contract["campos"].items():
            if rule["tipo_dato"] != "string":
                raise ValueError("Esta prueba solo expone campos string.")
            fields[field] = {
                key: rule[key]
                for key in ("marcador", "obligatorio", "max_caracteres", "max_lineas") if key in rule
            }
        types[kind] = {"pagina": contract["pagina"], "campos": fields}
    return Theme("poc_notas", Path(template), types)


def catalog_markdown(catalog):
    lines = ["# Catálogo experimental extraído del PPTX", "", "No editar: archivo derivado.", ""]
    for kind, contract in catalog["tipos"].items():
        lines += [f"## {kind}", "", contract["proposito"], "", "| Campo | Tipo | Obligatorio | Máximos |",
                  "| --- | --- | --- | --- |"]
        for field, rule in contract["campos"].items():
            required = "sí" if rule.get("obligatorio", False) else "no; omitido = vacío"
            limits = (
                f"{rule.get('max_caracteres', 'sin límite')} caracteres / "
                f"{rule.get('max_lineas', 'sin límite')} líneas"
            )
            lines.append(f"| {field} | {rule['tipo_dato']} | {required} | {limits} |")
        example = {"titulo": "Ejemplo de catálogo", "diapositivas": [contract["ejemplo"]]}
        lines += ["", "Ejemplo completo:", "", "```json", json.dumps(example, ensure_ascii=False, indent=2),
                  "```", "", "Restricciones: " + " ".join(contract["restricciones"]), ""]
    return "\n".join(lines)


def notes_state(template):
    result = []
    for slide in Presentation(str(template)).slides:
        if not slide.has_notes_slide:
            result.append({"texto": "", "relaciones": []})
            continue
        notes = slide.notes_slide
        frame = notes.notes_text_frame
        result.append({"texto": frame.text if frame else None,
                       "relaciones": sorted(r.reltype.rsplit("/", 1)[-1] for r in notes.part.rels.values())})
    return result


def sanitize_template(source, destination):
    """Drop ALL notes relationships on a disposable copy; return human text by page.

    Notes are detached rather than copying malformed notes parts. Human text is
    re-created with the public notes API AFTER generation. Formatting is not kept.
    """
    presentation = Presentation(str(source))
    human_notes = {}
    for page, slide in enumerate(presentation.slides, start=1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            raise ValueError(f"Página {page}: notas no textuales requieren revisión manual.")
        _, human_notes[page] = split_notes(frame.text)
        for rel in list(slide.part.rels.values()):
            if rel.reltype == RT.NOTES_SLIDE:
                slide.part.drop_rel(rel.rId)
    presentation.save(str(destination))
    return human_notes


def xml_contains(path, token):
    with zipfile.ZipFile(path) as package:
        return any(token.encode() in package.read(name) for name in package.namelist()
                   if name.endswith((".xml", ".rels")))


def run(output):
    theme = load_theme()  # Bootstrap ONLY: reuse the existing portada and its contract.
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    authored = output / "plantilla_con_metadatos.pptx"
    raw = output / "salida_sin_sanear_NO_DISTRIBUIR.pptx"
    sanitized = output / "plantilla_saneada_temporal.pptx"
    final = output / "presentacion_limpia.pptx"
    original_digest = hashlib.sha256(theme.template.read_bytes()).hexdigest()

    presentation = Presentation(str(theme.template))
    # Keep this historical pilot limited to one annotated type after production migration.
    for model_slide in presentation.slides:
        if model_slide.has_notes_slide:
            model_frame = model_slide.notes_slide.notes_text_frame
            if model_frame is not None:
                _, human = split_notes(model_frame.text)
                model_frame.text = human
    metadata = {
        "metadatos_version": 1, "tipo": "portada", "contrato_version": 1,
        "proposito": "Presentar el tema y orientar al alumnado al comienzo de la clase.",
        "campos": {field: {**rule, "tipo_dato": "string"} for field, rule in theme.types["portada"]["campos"].items()},
        "ejemplo": {"tipo": "portada", "titulo": "Herencia en Java", "subtitulo": "Reutilización y especialización"},
        "restricciones": ["Texto literal; sin Markdown.", "Diseño fijo; no hay ajuste automático por contenido."],
    }
    slide = presentation.slides[theme.types["portada"]["pagina"] - 1]
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        raise ValueError("La portada no tiene un cuadro de notas disponible.")
    frame.text = "Recordatorio docente: presentar el problema antes de la teoría.\n\n" + BEGIN + "\n" + json.dumps(
        metadata, ensure_ascii=False, indent=2
    ) + "\n" + END
    presentation.save(str(authored))

    catalog = extract_catalog(authored)
    extracted_theme = adapter_theme(catalog, authored)
    example = {"titulo": "Clase piloto", "diapositivas": [metadata["ejemplo"], {"tipo": "portada", "titulo": "Repaso"}]}
    config = adapt_content(example, extracted_theme)
    TemplateEngine(authored, render_rich_text=False).process(config, raw)

    human = sanitize_template(authored, sanitized)
    TemplateEngine(sanitized, render_rich_text=False).process(config, final)
    clean = Presentation(str(final))
    for generated_slide, spec in zip(clean.slides, config["slides"]):
        note = human.get(spec["src_page"], "")
        if note:
            generated_slide.notes_slide.notes_text_frame.text = note
    clean.save(str(final))

    report = {
        "python_pptx": version("python-pptx"), "pptx_slide_copier": version("pptx-slide-copier"),
        "notas_escritas_y_releidas": notes_state(authored)[0]["texto"].find(BEGIN) >= 0,
        "salida_sin_sanear_contiene_metadatos": xml_contains(raw, "PPTGEN_METADATA_JSON_V1"),
        "notas_salida_sin_sanear": notes_state(raw),
        "salida_limpia_contiene_metadatos": xml_contains(final, "PPTGEN_METADATA_JSON_V1"),
        "salida_limpia_contiene_contrato": xml_contains(final, "contrato_version"),
        "notas_salida_limpia": notes_state(final),
        "diapositivas_generadas": len(clean.slides),
        "subtitulo_opcional_vacio": any(
            shape.has_text_frame and shape.text == "" for shape in clean.slides[1].shapes
        ),
        "plantilla_original_intacta": hashlib.sha256(theme.template.read_bytes()).hexdigest() == original_digest,
    }
    assert report["notas_escritas_y_releidas"]
    assert report["salida_sin_sanear_contiene_metadatos"]
    assert not report["salida_limpia_contiene_metadatos"]
    assert not report["salida_limpia_contiene_contrato"]
    assert report["plantilla_original_intacta"] and report["subtitulo_opcional_vacio"]
    assert all(note["relaciones"] == ["notesMaster", "slide"] for note in report["notas_salida_limpia"])
    (output / "catalogo.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "catalogo_ia.md").write_text(catalog_markdown(catalog), encoding="utf-8")
    (output / "informe.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--salida", type=Path, default=Path("salida/poc_notas"))
    run(parser.parse_args().salida)
